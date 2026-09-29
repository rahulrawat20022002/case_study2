"""Rebuild the CS2 deck on the SRH template: template layouts, CS2 content, no sample slides."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree

TPL = "tpl2.pptx"
OUT = "CS2_Presentation_SRH.pptx"
FIG = "/home/claude/case_study2/report_figs/"
FOOTER = "SRH Hochschule Heidelberg – Case Study 2, Data Science and Analytics"

# SRH theme colours
ORANGE = RGBColor(0xDF, 0x48, 0x07)
GREY = RGBColor(0x57, 0x57, 0x56)     # dk1, body text
WARM = RGBColor(0xAA, 0xA3, 0x9D)     # lt2 warm grey
TEAL = RGBColor(0x35, 0xB4, 0xA0)
NAVY = RGBColor(0x02, 0x1E, 0x30)
TINT = RGBColor(0xF3, 0xF1, 0xEF)     # light warm tint for cards
ORANGE_TINT = RGBColor(0xFB, 0xE4, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation(TPL)
tslides = list(prs.slides)
T_TITLE, T_DIV, T_CONTENT, T_OLDCLOSE = tslides[0], tslides[1], tslides[2], tslides[20]
L_DIV, L_CONTENT, L_CLOSE = T_DIV.slide_layout, T_CONTENT.slide_layout, T_OLDCLOSE.slide_layout
FOOT_SRC = T_CONTENT  # cleaned template: dividers have no footer shapes left; placeholders inherit colour from their layout


def copy_footer(src, dst):
    """add_slide() does not clone date/footer/number placeholders; copy them from the template slide."""
    for sh in src.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx in (10, 11, 12):
            el = copy.deepcopy(sh._element)
            dst.shapes._spTree.append(el)
    set_footer(dst)


def set_footer(slide):
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == 11:
            p = sh.text_frame.paragraphs[0]
            runs = p.runs
            runs[0].text = FOOTER
            for r in runs[1:]:
                r._r.getparent().remove(r._r)


def ph(slide, idx):
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == idx:
            return sh


def remove_ph(slide, idx):
    sh = ph(slide, idx)
    if sh is not None:
        sh._element.getparent().remove(sh._element)


def divider(num, title):
    s = prs.slides.add_slide(L_DIV)
    copy_footer(FOOT_SRC, s)
    ph(s, 0).text_frame.text = title
    ph(s, 1).text_frame.text = num
    return s


def content(title, notes=None, keep_body=False):
    s = prs.slides.add_slide(L_CONTENT)
    copy_footer(FOOT_SRC, s)
    ph(s, 0).text_frame.text = title
    if not keep_body:
        remove_ph(s, 1)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def tb(slide, x, y, w, h, paras, size=14, color=GREY, bold=False, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, margin=0.0, space_after=4):
    """paras: str or list of str / (text, dict) tuples; dict may hold size, color, bold, italic."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if isinstance(paras, str):
        paras = [paras]
    for i, item in enumerate(paras):
        text, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opt.get("align", align)
        p.space_after = Pt(opt.get("space_after", space_after))
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", bold)
        f.italic = opt.get("italic", False)
        f.color.rgb = opt.get("color", color)
        if opt.get("bullet"):
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.25)))
            pPr.set("indent", str(-Inches(0.25)))
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "–")
    return box


def card(slide, x, y, w, h, fill=TINT):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def stat(slide, x, y, w, big, label, color=ORANGE, big_size=34):
    tb(slide, x, y, w, 0.65, big, size=big_size, color=color, bold=True)
    tb(slide, x, y + 0.68, w, 0.75, label, size=12, color=GREY)


def num_circle(slide, x, y, n, d=0.42, fill=ORANGE):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = WHITE


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def style_chart(chart, colors, legend=True, number_format='0.000', label_pos=XL_LABEL_POSITION.OUTSIDE_END,
                val_max=None, font=11, show_labels=True):
    chart.font.size = Pt(font)
    chart.font.color.rgb = GREY
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(font)
    plot = chart.plots[0]
    for s, c in zip(plot.series, colors):
        fmt = s.format
        if chart.chart_type in (XL_CHART_TYPE.LINE_MARKERS,):
            fmt.line.color.rgb = c
            fmt.line.width = Pt(2.5)
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = c
            s.marker.format.line.color.rgb = c
            s.smooth = False
        else:
            fmt.fill.solid()
            fmt.fill.fore_color.rgb = c
    if show_labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = number_format
        dl.number_format_is_linked = False
        dl.font.size = Pt(font - 1)
        dl.font.color.rgb = GREY
        if label_pos is not None:
            dl.position = label_pos
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE4, 0xE1, 0xDE)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(font - 1)
    if val_max is not None:
        va.maximum_scale = val_max
        va.minimum_scale = 0
    ca = chart.category_axis
    ca.tick_labels.font.size = Pt(font)
    ca.format.line.color.rgb = WARM
    ca.has_major_gridlines = False


# ---------------------------------------------------------------- 1. title (reuse template slide 1)
s = T_TITLE
title_ph = ph(s, 0)
title_ph.text_frame.text = "Faithful but wrong?"
sub = ph(s, 1)
lines = [
    ("Evaluating RAG-based high-risk classification under the EU AI Act against an authoritative answer key", True),
    ("Master of Science Data Science and Analytics – Case Study 2 – September 2026", False),
    ("Rahul Rawat  |  Supervisor: Prof. Dr. Swati Chandna", False),
]
tf = sub.text_frame
first_p = tf.paragraphs[0]
tmpl_rPr = copy.deepcopy(first_p.runs[0]._r.find(qn("a:rPr"))) if first_p.runs else None
for p in list(tf.paragraphs)[1:]:
    p._p.getparent().remove(p._p)
for r in list(first_p.runs):
    r._r.getparent().remove(r._r)
for br in first_p._p.findall(qn("a:br")):
    first_p._p.remove(br)
for i, (text, bold) in enumerate(lines):
    p = first_p if i == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = text
    r.font.bold = bold
# ---- remove the sample stock photo; the layout's orange field carries the slide
pic = ph(s, 21)
rid = pic._element.xpath(".//a:blip/@r:embed")
pic._element.getparent().remove(pic._element)
for r in rid:
    s.part.drop_rel(r)
s.background.fill.solid()
s.background.fill.fore_color.rgb = ORANGE
for shp in s.shapes:  # logo and arc are orange in the template: make them white on the orange field
    if shp.name in ("Logo flex.", "Teil des Logos"):
        shp.fill.solid()
        shp.fill.fore_color.rgb = WHITE
        if shp.name == "Teil des Logos":
            clr = shp.fill._xPr.find(qn("a:solidFill"))[0]
            etree.SubElement(clr, qn("a:alpha")).set("val", "18000")
sub.width = Inches(11.0)
s.notes_slide.notes_text_frame.text = ("Title. One sentence: this study asks whether the standard faithfulness metric for "
                                        "RAG systems tells you anything about whether a legal classification is actually correct.")

new = []  # slides in order (after the title)

# ---------------------------------------------------------------- 01 motivation
new.append(divider("01", "Motivation and task"))

s = content("Is grounded the same as right?",
            "Faithfulness checks grounding, not correctness. Crossing the two gives four buckets; the off-diagonal "
            "ones are the findings. The study is designed to give a real answer either way.")
new.append(s)
tb(s, 0.39, 1.45, 12.2, 0.5, "RAG compliance tools are usually judged by how grounded their answers are.", size=16)
for i, (head, body) in enumerate([
    ("Faithfulness (RAGAS)", "Share of the answer's claims that can be traced to the retrieved passages. Needs no answer key, which is why it is so widely used."),
    ("Correctness", "Does the predicted label match the official answer? Needs a ground truth, which legal tasks rarely have."),
]):
    y = 2.15 + i * 2.05
    card(s, 0.39, y, 5.6, 1.8)
    tb(s, 0.6, y + 0.2, 5.2, 0.4, head, size=18, bold=True, color=ORANGE)
    tb(s, 0.6, y + 0.7, 5.2, 1.0, body, size=14)
tb(s, 6.6, 2.15, 6.0, 0.4, "Two independent axes, four buckets", size=16, bold=True)
grid = [("Right and faithful", "expected", TINT, GREY), ("Right but unfaithful", "finding", ORANGE_TINT, ORANGE),
        ("Faithful but wrong", "headline finding", ORANGE, WHITE), ("Wrong and unfaithful", "expected", TINT, GREY)]
for k, (a, b, fill, col) in enumerate(grid):
    x = 6.6 + (k % 2) * 3.05
    y = 2.7 + (k // 2) * 1.62
    card(s, x, y, 2.9, 1.47, fill)
    tb(s, x + 0.15, y + 0.3, 2.6, 0.9, [(a, {"size": 15, "bold": True, "color": col}),
                                         (b, {"size": 12, "color": col, "italic": True})], align=PP_ALIGN.CENTER)

s = content("The task: Article 6 high-risk classification",
            "The legal test in four steps. The answer key comes from the Commission's own worked examples, which is "
            "what makes a correctness axis possible at all.")
new.append(s)
tb(s, 0.39, 1.45, 12.2, 0.4, "Annex III route of the EU AI Act, including the Article 6(3) filter", size=16)
steps = [("Intended purpose", "What does the system do, for whom, with which data?"),
         ("Annex III use case?", "Eight areas: biometrics, critical infrastructure, education, employment, essential services, law enforcement, migration, justice and democracy"),
         ("Article 6(3) filter", "Out again if only a narrow procedural or preparatory task, or improving completed human work. Profiling always stays high-risk."),
         ("Label", "high-risk or not-high-risk, with an explanation")]
w = 2.85
for i, (h, b) in enumerate(steps):
    x = 0.39 + i * (w + 0.27)
    card(s, x, 2.1, w, 2.75)
    num_circle(s, x + 0.2, 2.3, i + 1)
    tb(s, x + 0.2, 2.9, w - 0.4, 0.45, h, size=16, bold=True, color=ORANGE)
    tb(s, x + 0.2, 3.4, w - 0.4, 1.4, b, size=12.5)
card(s, 0.39, 5.15, 12.2, 1.35, ORANGE_TINT)
tb(s, 0.6, 5.3, 11.8, 1.1,
   [("Ground truth: ", {"bold": True, "size": 14, "color": ORANGE}),
    ("the European Commission's draft guidelines on Annex III (19 May 2026, frozen). Every worked example ends in an "
     "authoritative verdict, so each test input has one known correct answer. Article 6(1) product safety is out of scope.",
     {"size": 14})], space_after=0)

# ---------------------------------------------------------------- 02 data and method
new.append(divider("02", "Data and method"))

s = content("Test set: 215 inputs with official labels",
            "The source has no high-risk examples for law enforcement, migration and justice, so those areas test "
            "correct rejection only. Imbalance handled by reporting precision, recall and F1 next to accuracy.")
new.append(s)
tb(s, 0.39, 1.45, 8, 0.4, "Extracted from the worked examples; all eight Annex III areas covered", size=16)
cd = CategoryChartData()
cd.categories = ["Biometrics", "Critical infra.", "Education", "Employment", "Essential services",
                 "Law enforcement", "Migration", "Justice & democracy"]
cd.add_series("High-risk", (14, 8, 9, 14, 21, 0, 0, 0))
cd.add_series("Not high-risk", (18, 10, 19, 18, 22, 23, 23, 16))
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(0.39), Inches(2.0), Inches(8.2), Inches(4.7), cd).chart
style_chart(ch, [ORANGE, WARM], number_format='0;;;', label_pos=XL_LABEL_POSITION.CENTER)
ch.plots[0].data_labels.font.color.rgb = WHITE
ch.plots[0].gap_width = 60
ch.category_axis.reverse_order = True
stat(s, 9.2, 2.1, 3.4, "215", "test inputs, one gold label each")
stat(s, 9.2, 3.6, 3.4, "66 / 149", "high-risk / not high-risk (31% / 69%)")
stat(s, 9.2, 5.1, 3.4, "41", "Article 6(3) filter cases, all not high-risk")

s = content("Evaluation framework",
            "Walk the diagram top to bottom: source split, shared retrieval, three configurations, evaluation "
            "harness, four buckets. Everything is held fixed except the prompt. A local open generator replaced the "
            "planned paid API model: it can be pinned, rerun and costs nothing, and it proved fully reliable.")
new.append(s)
picture(s, FIG + "arch_detailed.png", 0.39, 1.4, h=5.35)
pts = [("Source split", "Worked examples become the test set; only normative text (400 chunks) goes into the retrieval corpus, so labels cannot leak through retrieval."),
       ("Three systems, one generator", "phi-4 (14B) for all: plain LLM with no retrieval; standard RAG with top 5 passages (hybrid BM25 + dense) and a generic prompt; agent with the same passages and a prompt that walks the legal test step by step."),
       ("Two independent axes", "Judge (Qwen2.5-72B, separate model family) scores RAGAS faithfulness; correctness is a plain lookup against the Commission label."),
       ("Reproducible", "Local models on the university GPU, pinned revisions, greedy decoding. 645 generations, 0 parse failures, 0 errors.")]
for i, (h, b) in enumerate(pts):
    y = 1.5 + i * 1.3
    num_circle(s, 5.0, y, i + 1)
    tb(s, 5.6, y - 0.02, 7.0, 0.4, h, size=16, bold=True, color=ORANGE)
    tb(s, 5.6, y + 0.4, 7.0, 0.85, b, size=13)

s = content("A label leak found and fixed",
            "Plain model: 0.94 on leaky rows, 0.67 on clean rows. Retrieval systems barely affected. Fixed by "
            "redaction and a full rerun; disclosed in the report.")
new.append(s)
tb(s, 0.39, 1.45, 12, 0.4, "60% of inputs contained the Commission's own verdict (first run, before redaction)", size=16)
cd = CategoryChartData()
cd.categories = ["Plain LLM", "Standard RAG", "Agent"]
cd.add_series("129 inputs stating the verdict", (0.938, 0.922, 0.961))
cd.add_series("86 clean inputs", (0.674, 0.953, 0.953))
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.39), Inches(2.0), Inches(6.6), Inches(4.7), cd).chart
style_chart(ch, [ORANGE, WARM], val_max=1.1)
ch.value_axis.tick_labels.number_format = '0.0'
ch.value_axis.tick_labels.number_format_is_linked = False
ch.value_axis.major_unit = 0.2
tb(s, 7.4, 2.1, 5.2, 0.4, "What happened", size=16, bold=True, color=ORANGE)
tb(s, 7.4, 2.55, 5.2, 1.3, "The extractor kept each example's closing verdict, e.g. \"It therefore falls outside the use "
   "case of point 4(a)\". The plain model was reading the answer.", size=13)
tb(s, 7.4, 3.85, 5.2, 0.4, "How it was fixed", size=16, bold=True, color=ORANGE)
tb(s, 7.4, 4.3, 5.2, 2.4, [
    ("Kept what the system does; removed every legal conclusion and citation", {"bullet": True}),
    ("96 rows automatic, 31 hand-edited, 46 reviewed by the author", {"bullet": True}),
    ("IDs and gold labels unchanged; all systems rerun", {"bullet": True}),
    ("Mean input length 557 to 392 characters", {"bullet": True})], size=13)

# ---------------------------------------------------------------- 03 results
new.append(divider("03", "Results"))

s = content("Result 1: retrieval helps",
            "Grounding in the guidelines removes most false alarms. The step-by-step agent is not better than a "
            "generic RAG prompt. The leak had even flipped the ranking of the two.")
new.append(s)
tb(s, 0.39, 1.45, 12, 0.4, "…but the structured prompt does not. Accuracy on 215 inputs, before and after redaction", size=16)
cd = CategoryChartData()
cd.categories = ["Plain LLM", "Standard RAG", "Agent"]
cd.add_series("Original inputs", (0.833, 0.935, 0.958))
cd.add_series("Redacted inputs", (0.688, 0.879, 0.856))
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.39), Inches(2.0), Inches(7.2), Inches(4.7), cd).chart
style_chart(ch, [WARM, ORANGE], val_max=1.1)
ch.value_axis.tick_labels.number_format = '0.0'
ch.value_axis.tick_labels.number_format_is_linked = False
ch.value_axis.major_unit = 0.2
stat(s, 8.1, 2.1, 4.5, "+19 pts", "standard RAG vs plain LLM on redacted inputs (McNemar p < 0.001)")
stat(s, 8.1, 3.6, 4.5, "60 vs 17", "false high-risk calls: plain LLM vs standard RAG")
stat(s, 8.1, 5.1, 4.5, "p = 0.30", "agent vs standard RAG: no significant difference (0.856 vs 0.879)")

s = content("Where the systems struggle",
            "Two patterns: critical infrastructure is hard for everyone because of the safety-component test; the "
            "plain model over-calls in law enforcement.")
new.append(s)
tb(s, 0.39, 1.45, 12, 0.4, "Accuracy by Annex III area and error types (redacted inputs)", size=16)
picture(s, FIG + "fig_area_heatmap.png", 0.39, 2.0, w=8.0)
picture(s, FIG + "fig_confusion.png", 0.39, 5.0, w=5.0)
card(s, 8.8, 2.0, 3.8, 2.1)
tb(s, 9.0, 2.15, 3.4, 0.4, "Critical infrastructure", size=15, bold=True, color=ORANGE)
tb(s, 9.0, 2.6, 3.4, 1.4, "Hardest for both RAG systems (0.61, 0.56). The test is safety component, not sector.", size=13)
card(s, 8.8, 4.35, 3.8, 2.1)
tb(s, 9.0, 4.5, 3.4, 0.4, "Law enforcement", size=15, bold=True, color=ORANGE)
tb(s, 9.0, 4.95, 3.4, 1.4, "Plain LLM collapses to 0.26: it calls most police uses high-risk. RAG fixes this (0.87+).", size=13)

s = content("Result 2: faithfulness misses wrong labels",
            "AUC is the chance that a random right answer scores higher than a random wrong one. 0.5 is a coin "
            "flip. For the agent the confidence interval includes 0.5.")
new.append(s)
tb(s, 0.39, 1.45, 12, 0.4, "RAGAS faithfulness of every answer, split by whether the label was correct", size=16)
picture(s, FIG + "fig3_faithfulness.png", 0.39, 2.0, w=8.6)
stat(s, 9.4, 2.1, 3.3, "0.465", "agent AUC (95% CI 0.36 to 0.57): chance level")
stat(s, 9.4, 3.6, 3.3, "0.49 vs 0.45", "agent: wrong answers score slightly higher than right ones", big_size=30)
stat(s, 9.4, 5.1, 3.3, "0.615", "standard RAG AUC: weak signal only", color=GREY)

s = content("Both off-diagonal buckets are populated",
            "Faithful but wrong: 17 agent, 14 standard RAG. Right but unfaithful: about half of correct answers.")
new.append(s)
tb(s, 0.39, 1.45, 12, 0.4, "Four-bucket matrix at a faithful cut of 0.5", size=16)
picture(s, FIG + "fig_buckets.png", 0.39, 2.0, w=8.6)
stat(s, 9.4, 2.1, 3.3, "17 of 31", "agent errors that pass as faithful (55%)")
stat(s, 9.4, 3.6, 3.3, "93 of 184", "correct agent answers scored as unfaithful", color=GREY)
tb(s, 0.39, 5.95, 8.6, 0.8, "Part of the right-but-unfaithful effect is mechanical: sentences that restate the input are "
   "not in the passages, so the judge marks them unsupported. That is how RAGAS is defined.", size=12, color=GREY)

s = content("Not an artefact of the cut-off",
            "Counts: agent 22, 17, 7, 1 of 31. Standard RAG 21, 14, 1, 0 of 26.")
new.append(s)
tb(s, 0.39, 1.45, 12, 0.4, "Share of wrong answers that still pass as faithful as the threshold rises", size=16)
cd = CategoryChartData()
cd.categories = ["0.3", "0.5", "0.7", "0.9"]
cd.add_series("Standard RAG", (0.808, 0.538, 0.038, 0.0))
cd.add_series("Agent", (0.710, 0.548, 0.226, 0.032))
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.39), Inches(2.0), Inches(7.6), Inches(4.7), cd).chart
style_chart(ch, [WARM, ORANGE], number_format='0%', label_pos=None, val_max=1.0, show_labels=False)
agent = ch.plots[0].series[1]
dls = agent.data_labels
dls.show_value = True
dls.number_format = '0%'
dls.number_format_is_linked = False
dls.position = XL_LABEL_POSITION.ABOVE
dls.font.size = Pt(11)
dls.font.color.rgb = ORANGE
dls.font.bold = True
ch.value_axis.tick_labels.number_format = '0%'
ch.value_axis.tick_labels.number_format_is_linked = False
ch.value_axis.major_unit = 0.2
ch.category_axis.has_title = True
ch.category_axis.axis_title.text_frame.text = "Faithful cut (score counted as faithful at or above)"
ch.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
ch.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.bold = False
stat(s, 8.5, 2.1, 4.1, "7", "agent errors still count as faithful at a strict cut of 0.7")
stat(s, 8.5, 3.6, 4.1, "1.00", "perfect faithfulness score on a wrong answer (input 207)")
tb(s, 8.5, 5.15, 4.1, 1.3, "Threshold-free measures (AUC, correlation) are reported alongside, so the finding does not "
   "rest on one cut.", size=12, color=GREY)

s = content("Why the wrong answers look grounded",
            "The model had the right rule but applied it wrongly. The judge accepts the conclusion because the "
            "passages state the general rule correctly. Faithfulness never checks the application step.")
new.append(s)
tb(s, 0.39, 1.45, 12, 0.4, "All 17 faithful-but-wrong agent answers read and categorised", size=16)
cd = CategoryChartData()
cd.categories = ["A. Sector taken as the use case", "B. 6(3) exception not applied", "C. Profiling rule misused",
                 "D. Condition misread / fact invented", "E. Too cautious on short input"]
cd.add_series("Answers", (6, 4, 3, 2, 2))
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.39), Inches(1.95), Inches(6.4), Inches(3.3), cd).chart
style_chart(ch, [ORANGE], legend=False, number_format='0', val_max=7)
ch.has_title = False
ch.plots[0].gap_width = 70
ch.category_axis.reverse_order = True
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
stat(s, 0.39, 5.35, 3.0, "14 / 17", "had the correct guideline section retrieved")
stat(s, 3.6, 5.35, 3.2, "21 / 31", "verdict statements judged 'supported'")
card(s, 7.2, 1.95, 5.4, 4.75, ORANGE_TINT)
tb(s, 7.45, 2.1, 4.9, 4.5, [
    ("Case: input 207", {"size": 16, "bold": True, "color": ORANGE, "space_after": 8}),
    ("A tool that pre-classifies incoming court cases by type.", {"size": 13, "space_after": 10}),
    ("Agent: ", {"size": 13, "bold": True, "space_after": 0}),
    ("places it under point 8(a), then asserts no Article 6(3) condition applies. Label: high-risk.", {"size": 13, "space_after": 10}),
    ("Commission: ", {"size": 13, "bold": True, "space_after": 0}),
    ("narrow procedural / preparatory task. Not high-risk.", {"size": 13, "space_after": 10}),
    ("Faithfulness: 6 of 6 statements supported = 1.00", {"size": 14, "bold": True, "color": ORANGE}),
])

# ---------------------------------------------------------------- 04 conclusion
new.append(divider("04", "Conclusion"))

s = content("Key findings", "Four takeaways, most important first.")
new.append(s)
kf = [("Faithfulness is not correctness", "For the agent, faithfulness is at chance level for spotting wrong labels (AUC 0.465). 17 of 31 errors pass as faithful."),
      ("Errors are application errors", "The right passage was retrieved in 14 of 17 cases; the model misapplied the rule, which RAGAS cannot see."),
      ("Retrieval matters, prompt engineering did not", "RAG adds about 19 points over the plain model; the step-by-step agent is not better than a generic RAG prompt."),
      ("Benchmarks from worked examples can leak", "The verdict text inflated accuracy by up to 15 points and flipped the ranking of the two best systems.")]
for i, (h, b) in enumerate(kf):
    x = 0.39 + (i % 2) * 6.2
    y = 1.6 + (i // 2) * 2.55
    card(s, x, y, 5.95, 2.3, ORANGE_TINT if i == 0 else TINT)
    num_circle(s, x + 0.25, y + 0.25, i + 1)
    tb(s, x + 0.85, y + 0.25, 4.9, 0.5, h, size=16, bold=True, color=ORANGE)
    tb(s, x + 0.85, y + 0.85, 4.9, 1.35, b, size=13)

s = content("Limitations and future work", "Be upfront about the limits; each maps to a concrete next step.")
new.append(s)
card(s, 0.39, 1.6, 5.95, 4.2)
card(s, 6.59, 1.6, 5.95, 4.2, ORANGE_TINT)
tb(s, 0.64, 1.8, 5.5, 0.45, "Limitations", size=18, bold=True, color=GREY)
tb(s, 0.64, 2.4, 5.5, 4.1, [(t, {"bullet": True, "space_after": 8}) for t in [
    "One generator and one judge; results may depend on these models",
    "Single annotator for redaction review and error categories",
    "About 20 inputs still describe the legal test without stating the label",
    "26 and 31 wrong answers: wide confidence intervals",
    "Commission examples are cleaner than real product descriptions"]], size=13.5)
tb(s, 6.84, 1.8, 5.5, 0.45, "Future work", size=18, bold=True, color=ORANGE)
tb(s, 6.84, 2.4, 5.5, 4.1, [(t, {"bullet": True, "space_after": 8}) for t in [
    "A second, stronger generator, including a closed commercial model, on the same redacted inputs",
    "A second annotator and inter-annotator agreement",
    "Realistic product descriptions labelled against the guidelines",
    "A correctness-aware check: does the conclusion follow from the retrieved rule and the stated facts?",
    "Extension towards a thesis or paper"]], size=13.5)

# ---------------------------------------------------------------- closing (fresh slide on the statement layout)
s = prs.slides.add_slide(L_CLOSE)
copy_footer(FOOT_SRC, s)
T_CLOSE = s
t = ph(s, 0).text_frame
lines = ["A high faithfulness score says almost nothing about whether a legal classification is right.",
         "",
         "Thank you. Questions?",
         "Danke für Ihre Aufmerksamkeit!"]
p0 = t.paragraphs[0]
for p in list(t.paragraphs)[1:]:
    p._p.getparent().remove(p._p)
for r in list(p0.runs):
    r._r.getparent().remove(r._r)
for br in p0._p.findall(qn("a:br")):
    p0._p.remove(br)
for i, line in enumerate(lines):
    p = p0 if i == 0 else t.add_paragraph()
    r = p.add_run()
    r.text = line
    if i == 0:
        r.font.size = Pt(26)
tb(s, 0.39, 6.35, 8, 0.4, "github.com/rahulrawat20022002/case_study2", size=12, color=WHITE)
s.notes_slide.notes_text_frame.text = ("Close on the one-line takeaway and the practical recommendation: faithfulness is "
                                        "necessary but not sufficient; wherever a ground truth exists, report a correctness "
                                        "check next to it.")

# ---------------------------------------------------------------- order: title, new slides, closing; drop samples
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
by_part = {prs.slides.part.related_part(i.get(qn("r:id"))): i for i in ids}
keep = [T_TITLE] + new + [T_CLOSE]
keep_ids = [by_part[sl.part] for sl in keep]
for el in ids:
    if el not in keep_ids:
        prs.part.drop_rel(el.get(qn("r:id")))
    sldIdLst.remove(el)
for el in keep_ids:
    sldIdLst.append(el)
prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")


# ---------------------------------------------------------------- post-process: drop the template's embedded
# SRH font SUBSETS (they only hold the glyphs of the sample text, so any other letter renders as garbage in
# PowerPoint) and set every SRH typeface to Arial so the deck renders the same on any machine.
import re, zipfile, shutil, os
def fix_fonts(path):
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        pres = zin.read("ppt/presentation.xml").decode("utf-8")
        m = re.search(r"<p:embeddedFontLst>.*?</p:embeddedFontLst>", pres)
        rids = re.findall(r'r:id="(rId\d+)"', m.group(0)) if m else []
        for item in zin.infolist():
            name = item.filename
            if name.startswith("ppt/fonts/"):
                continue
            data = zin.read(name)
            if name.endswith(".xml") or name.endswith(".rels"):
                s = data.decode("utf-8")
                if name == "ppt/presentation.xml" and m:
                    s = s.replace(m.group(0), "")
                    s = re.sub(r'\s(embedTrueTypeFonts|saveSubsetFonts)="[^"]*"', "", s)
                if name == "ppt/_rels/presentation.xml.rels":
                    for rid in rids:
                        s = re.sub(r'<Relationship [^>]*Id="%s"[^>]*/>' % rid, "", s)
                s = re.sub(r'typeface="SRH[^"]*"', 'typeface="Arial"', s)
                s = re.sub(r'(<a:latin typeface="Arial")\s+panose="[^"]*"', r"\1", s)
                data = s.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)

fix_fonts(OUT)
print("fonts fixed")
