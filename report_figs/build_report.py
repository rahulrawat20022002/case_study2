"""Build CS2_Report.docx on top of the SRH report template (styles, page setup,
cover layout and table style reused from the template file)."""
import copy, json
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Pt, RGBColor, Inches, Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEMPLATE = "/tmp/claude-0/-home-claude/3f0a219a-3c0a-53e0-8360-748fb7d57569/scratchpad/sar.docx"
OUT = "/home/claude/case_study2/report_figs/CS2_Report.docx"
FIG = "/home/claude/case_study2/report_figs/"
M = json.load(open("/tmp/claude-0/-home-claude/3f0a219a-3c0a-53e0-8360-748fb7d57569/scratchpad/metrics_all.json"))

doc = docx.Document(TEMPLATE)
body = doc.element.body
sectPr = body.find(qn("w:sectPr"))
for el in list(body):
    if el is not sectPr:
        body.remove(el)

RED = RGBColor(0xC0, 0x39, 0x2B); BLUE = RGBColor(0x2E, 0x74, 0xB5); GREY = RGBColor(0x55, 0x55, 0x55)
TABLE_STYLE = "Light Grid Accent 1"


def P(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11, bold=False, italic=False, color=None, style=None, space_after=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if text:
        add_runs(p, text, size=size, bold=bold, italic=italic, color=color)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def add_runs(p, text, size=11, bold=False, italic=False, color=None):
    """**bold** segments supported inside text."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.font.size = Pt(size)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if color is not None:
            r.font.color.rgb = color
    return p


HEADINGS = []


def H1(t):
    HEADINGS.append((1, t)); return doc.add_paragraph(t, style="Heading 1")


def H2(t):
    HEADINGS.append((2, t)); return doc.add_paragraph(t, style="Heading 2")


def H3(t):
    return doc.add_paragraph(t, style="Heading 3")


def caption(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = BLUE
    return p


def figure(path, width_in, cap):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width_in))
    caption(cap)


def table(rows, widths, cap, header=True, bold_first_col=False):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = doc.styles[TABLE_STYLE]
    t.alignment = 1
    t.autofit = False
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[j] * 1440)))
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for row in t.rows:
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); row._tr.get_or_add_trPr().append(cs)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Inches(widths[j])
            c.paragraphs[0].text = ""
            r = c.paragraphs[0].add_run(str(val)); r.font.size = Pt(9)
            c.paragraphs[0].paragraph_format.keep_with_next = True
            if i == 0 and header:
                r.bold = True
    caption(cap)
    return t


def bullet_list(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet") if "List Bullet" in [s.name for s in doc.styles] else doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, it, size=11)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


TOC_ANCHOR = []


def toc():
    TOC_ANCHOR.append(doc.add_paragraph())


def highlight_note(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text); r.font.size = Pt(11); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


f3 = lambda x: f"{x:.3f}"
f2 = lambda x: f"{x:.2f}"
L, RD = M["leaky"], M["redacted"]
B1, B2, AG = "baseline1_plain_llm", "baseline2_standard_rag", "agent_structured"

# ------------------------------------------------------------------ cover
P("srh", align=WD_ALIGN_PARAGRAPH.CENTER, size=36, bold=True, color=RED)
P("SRH Hochschule Heidelberg", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, color=GREY)
P("")
P("FAITHFUL BUT WRONG? EVALUATING RETRIEVAL-AUGMENTED CLASSIFICATION OF HIGH-RISK AI SYSTEMS UNDER THE EU AI ACT AGAINST AN AUTHORITATIVE ANSWER KEY",
  align=WD_ALIGN_PARAGRAPH.CENTER, size=18, bold=True, color=BLUE)
P("")
for t, c in [("BY", None), ("Rahul Rawat", RED), ("[Matriculation number]", RED), ("M.Sc. Data Science and Analytics", None),
             ("SRH Hochschule Heidelberg", None), ("Case Study 2", None), ("September 2026", None)]:
    P(t, align=WD_ALIGN_PARAGRAPH.CENTER, size=11, color=c)
P("")
P("Supervisor: Prof. Dr. Swati Chandna", align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
page_break()

# ------------------------------------------------------------------ abstract
H1("Abstract")
P("Retrieval-augmented generation (RAG) systems are increasingly proposed as compliance assistants, and they are usually "
  "evaluated with automated faithfulness metrics that check whether an answer is grounded in the retrieved text. This case "
  "study asks whether such a metric says anything about whether the legal answer is actually right. The task is the "
  "high-risk classification of AI systems under Article 6(2) and Annex III of the EU AI Act, including the Article 6(3) "
  "filter. The answer key is taken from the worked examples in the European Commission's May 2026 draft guidelines, which "
  "gives 215 test inputs with authoritative labels. Three systems built on one shared retrieval pipeline are compared: a "
  "plain language model without retrieval, a standard RAG system, and a structured agent whose prompt walks through the "
  "legal test step by step. Answers are scored on two independent axes: correctness against the Commission label, and "
  "RAGAS faithfulness computed by a separate 72B judge model. During the study, 129 of the 215 inputs were found to "
  "contain the Commission's verdict text; these were redacted and all systems were run again. On the redacted inputs, "
  "retrieval clearly helps (accuracy 0.688 without retrieval against 0.879 and 0.856 with it), but the structured agent "
  "is not better than standard RAG. For the agent, faithfulness does not separate right from wrong answers (AUC 0.465, "
  "95% CI 0.36 to 0.57), and 17 of its 31 wrong answers still pass as faithful. Error analysis shows that these are "
  "mostly errors in applying a correctly retrieved rule, which the metric is not designed to detect. The practical "
  "conclusion is that faithfulness is not a proxy for legal correctness and needs to be reported next to a correctness "
  "check whenever a ground truth exists.")
P("**Index Terms**: EU AI Act, high-risk classification, retrieval-augmented generation, RAGAS, faithfulness, "
  "LLM evaluation, LLM-as-a-judge, legal NLP, label leakage.", size=11)
page_break()

H1("Contents")
toc()
page_break()

# ------------------------------------------------------------------ 1
H1("1. Introduction")
P("The EU AI Act (Regulation (EU) 2024/1689) [1] sorts AI systems by risk, and the most consequential step for a "
  "provider is deciding whether a system is high-risk. For the Annex III route this means checking whether the system's "
  "intended purpose falls within one of eight listed areas, such as employment, education or access to essential "
  "services, and then whether one of the Article 6(3) conditions lets it out again. The European Commission's draft "
  "guidelines of May 2026 [2] explain this test and illustrate it with many worked examples, each ending in a clear "
  "verdict.")
P("Tools that answer such questions by retrieving passages from legal text and letting a language model write the "
  "answer are an obvious fit, and they are commonly evaluated with reference-free metrics. The most widely used of these "
  "is faithfulness as defined in RAGAS [3]: the share of claims in an answer that can be inferred from the retrieved "
  "context. Faithfulness is attractive because it needs no answer key. By design, however, it only checks whether the "
  "explanation stays close to the retrieved text. It does not check whether the conclusion is legally correct.")
P("This case study measures how far apart those two things are on a task where the correct answer is known. The "
  "compliance system is the system under test, not the contribution; the contribution is the evaluation. The central "
  "question is whether correctness and faithfulness come apart, and in particular whether there are answers that are "
  "**faithful but wrong** (well grounded, legally incorrect) and answers that are **right but unfaithful** (correct "
  "label, reasoning not traceable to the retrieved text). The study was set up to give a real answer in either "
  "direction.")
P("The report is organised as follows. Section 2 gives background. Section 3 describes the task, the ground truth and "
  "the test set, including a label leak found during the study and how it was fixed. Section 4 describes the three "
  "systems and the evaluation harness. Section 5 reports the results and an error analysis, Section 6 discusses them, "
  "and Section 7 concludes with limitations and future work.")

# ------------------------------------------------------------------ 2
H1("2. Background and Related Work")
P("**High-risk classification under the AI Act.** Article 6 defines two routes to high-risk status. The Article 6(1) "
  "route covers AI systems that are safety components of products under Union harmonisation legislation listed in "
  "Annex I; it depends on a separate body of product law and is out of scope here. The Article 6(2) route covers the use "
  "cases listed in Annex III. Article 6(3) then acts as a filter: an Annex III system is not high-risk if it only "
  "performs a narrow procedural task, improves the result of a completed human activity, detects decision patterns "
  "without replacing human assessment, or performs a preparatory task, unless it profiles natural persons, in which case "
  "it is always high-risk. The Commission guidelines [2] turn this into practical guidance and are the only source of "
  "authoritative labels used in this study.")
P("**Evaluating RAG systems.** RAG [4] combines a retriever with a generator so that answers can draw on an external "
  "corpus. RAGAS [3] proposes reference-free metrics for such systems, among them faithfulness, computed in two steps: an "
  "LLM breaks the answer into atomic statements, and an LLM then decides for each statement whether it can be inferred "
  "from the retrieved context. ARES [5] trains lightweight judges for context relevance, answer faithfulness and answer "
  "relevance and calibrates them with a small set of human labels. Both frameworks evaluate grounding and relevance; "
  "neither is built to check a domain-specific correct answer, which is the gap this study measures.")
P("**Legal NLP benchmarks.** LexGLUE [6] collects legal classification tasks with gold labels and shows that legal text "
  "remains hard for general models. It evaluates label accuracy but not the grounding of explanations, which is the "
  "other half of the picture examined here.")
P("**Other frameworks.** The OECD framework for classifying AI systems [7] and the NIST AI Risk Management Framework [8] "
  "offer broader ways of describing and managing AI risk. They do not produce EU high-risk labels and are therefore used "
  "only as context; all labels in this study come from the Commission guidelines.")
highlight_note("[Note for Rahul, remove before sending: add Wallat et al. (2025), 'Correctness is not Faithfulness in RAG "
               "Attributions', here once you have read it in full. Position this study as the legal-domain case with an "
               "authoritative answer key, not as the first to observe the gap.]")

# ------------------------------------------------------------------ 3
H1("3. Task and Data")
H2("3.1 Task Definition")
P("Each test input is a short description of an AI system and its intended purpose. The system under test must return "
  "one of two labels, high-risk or not-high-risk, under the Article 6(2) and Annex III route including the Article 6(3) "
  "filter, together with an explanation. The label is scored for correctness; the explanation is scored for "
  "faithfulness. The Article 6(1) product-safety route is excluded because it has no clean ground truth in the source.")

H2("3.2 Ground Truth")
P("The ground truth is the European Commission's draft guidelines on the classification of high-risk AI systems for "
  "Annex III, May 2026 version [2]. The document was downloaded once and frozen; later revisions were deliberately not "
  "followed during the project. Its worked examples each describe a concrete system and state whether it is high-risk, "
  "which gives one known correct answer per input. Without such an answer key the claim that a metric missed a legal "
  "error could not be made.")

H2("3.3 Test Set Construction")
P("A Python parser extracted the worked examples from the text layer of the PDF, handling page breaks, footnotes and "
  "numbering. An automatic sanity check flagged suspected truncations and length outliers; 21 rows were repaired by "
  "stripping a heading of the next example that the parser had concatenated, and 4 rows were dropped as merged or "
  "truncated. No label was changed. The locked test set has 215 inputs covering all eight Annex III areas: 66 high-risk "
  "and 149 not-high-risk (31% and 69%), with 41 inputs that are Article 6(3) filter cases (Table 1). Because of the "
  "imbalance, precision, recall and F1 for the high-risk class are reported next to accuracy.")
table([["Annex III area", "Inputs", "High-risk", "Not high-risk"],
       ["1 Biometrics", 32, 14, 18], ["2 Critical infrastructure", 18, 8, 10], ["3 Education", 28, 9, 19],
       ["4 Employment", 32, 14, 18], ["5 Essential services", 43, 21, 22], ["6 Law enforcement", 23, 0, 23],
       ["7 Migration and border control", 23, 0, 23], ["8 Justice and democracy", 16, 0, 16], ["Total", 215, 66, 149]],
      [2.6, 0.9, 0.9, 1.1], "Table 1. Test set by Annex III area. 41 of the 215 inputs are Article 6(3) filter cases.")
figure(FIG + "fig_dataset.png", 5.6, "Figure 1. Composition of the test set by Annex III area: high-risk inputs, Article 6(3) filter cases and other not-high-risk inputs.")
P("The source has no high-risk worked examples for law enforcement, migration or justice and democracy (Figure 1); every example "
  "there is either a filter case or out of scope. The evaluation therefore tests high-risk detection in five areas and "
  "correct rejection in all eight. This is a property of the source and is reported as such.")

H2("3.4 Label Leakage Found and Fixed")
P("The first full run produced suspiciously high accuracy, and inspecting the inputs showed why: the Commission's "
  "examples usually end with the verdict and its legal reasoning (for example 'It therefore falls outside the use case "
  "of point 4(a) of Annex III'), and the extractor had kept this text as part of the system description. A keyword scan "
  "flagged 129 of the 215 inputs (60%) as stating or strongly implying the answer, including 36 of the 41 filter "
  "cases. The effect was large for the model without retrieval, which reached 0.938 accuracy on the flagged inputs but "
  "only 0.674 on the 86 clean ones, and small for the two retrieval systems (0.922 and 0.961 flagged, 0.953 and 0.953 "
  "clean).")
figure(FIG + "fig_leak.png", 5.0, "Figure 2. Accuracy in the first run on the 129 inputs whose text stated the verdict and on the 86 clean inputs. Only the model without retrieval depends heavily on the leak.")
P("The inputs were therefore redacted and all systems were run again. The rule was to keep what the system does, who "
  "uses it, what data it uses, what its output is used for and how humans are involved, and to remove every legal "
  "conclusion and every citation of an article, annex point, recital or section. IDs and gold labels were not touched, "
  "and the model never sees the label field. Redaction was done in three layers: an automatic filter that drops "
  "sentences containing legal vocabulary, hand edits where the verdict sat inside a descriptive sentence, and a full "
  "manual review of the hand-edited rows by the author (Table 2). In that review the author restored some wording that "
  "describes the legal test without stating the label (for example that a system has no relevant impact on a decision), "
  "while every explicit verdict was removed. This residual wording is disclosed as a limitation in Section 7. The "
  "redaction is fully scripted and the reviewed text is stored with the dataset, so the redacted set can be rebuilt "
  "exactly.")
table([["Handling", "Inputs"], ["Unchanged (no verdict text)", 42], ["Automatic sentence removal", 96],
       ["Hand-edited clause removal", 31], ["Hand-edited, then reviewed and revised by the author", 46], ["Total", 215]],
      [3.8, 1.0], "Table 2. Redaction of the test inputs. Mean description length fell from 557 to 392 characters.")
P("All results in Section 5 use the redacted inputs unless stated otherwise. The first run is kept for comparison only.")

# ------------------------------------------------------------------ 4
H1("4. Methodology")
P("Figure 3 shows the whole pipeline. One retrieval pipeline and one generator are shared by all three systems, so the "
  "only thing that changes between them is the prompt. The evaluation harness is separate from the systems and reads "
  "only their output files.")
figure(FIG + "arch_detailed.png", 5.9, "Figure 3. End-to-end evaluation framework: the source split that keeps worked examples out of the retrieval corpus, the shared retrieval pipeline, the three configurations with one generator, the evaluation harness with its two independent axes, and the four outcome buckets.")

H2("4.1 Retrieval Corpus and Index")
P("The retrieval corpus is built from the same guidelines document, but only from its normative text. Chunking is "
  "structural: one chunk per heading and numbered paragraph, with footnotes attached to the paragraph they annotate. The "
  "worked examples and verdict lists are held out, because they are the test set and would otherwise leak labels into "
  "retrieval, and the Annex I section is excluded as out of scope. The corpus has 400 chunks. They are embedded with "
  "BAAI/bge-base-en-v1.5 [9] and stored in a persistent Chroma index, and a BM25 [10] index is built in memory. "
  "Retrieval is hybrid with equal weighting (alpha 0.5, not tuned during evaluation). The top 5 passages go to the "
  "generator and the top 10 are logged for analysis. An earlier version of the corpus builder mistook footnote dates "
  "for headings and silently lost 39 chunks, including the whole recruitment section; this was caught by the retrieval "
  "evaluation and fixed before any generation run.")

H2("4.2 The Three Systems Under Test")
bullet_list([
    "**Baseline 1, plain LLM.** The description only, no retrieved text. The model returns a label and a short reason. "
    "It has no retrieved context, so it has no faithfulness score by design.",
    "**Baseline 2, standard RAG.** The description plus the top 5 passages and a generic instruction to classify the "
    "system and ground the explanation in the passages.",
    "**Agent, structured RAG.** The same passages, with a prompt that walks through the legal test: identify the "
    "intended purpose, check it against the eight Annex III areas, apply the Article 6(3) conditions, note that "
    "profiling always makes a system high-risk, and state which point and condition apply with reference to the "
    "passages.",
])
P("All three return strict JSON with a label and an explanation, which is parsed deterministically. The retrieval "
  "components were reused from an earlier RAG project of the author; its other agents (debate, verification, "
  "evaluation, memory, planning, guardrails, summarisation and topic modelling) were removed because they would either "
  "interfere with the faithfulness measurement or make runs non-reproducible.")

H2("4.3 Generator")
P("The generator is microsoft/phi-4 [11], a 14-billion-parameter open instruction model, served locally with vLLM "
  "[12] on the university GPU server with greedy decoding (temperature 0, top-p 1) and a pinned model revision. The "
  "original plan used a commercial API model. A local open model was chosen instead because it can be pinned to an "
  "exact version and rerun at any time, it gives full control over decoding and output length, and it costs nothing. "
  "Free hosted models tried during development could not be pinned and produced empty or truncated answers under their "
  "token limits; with the local setup, none of the 645 generations failed to parse, came back empty or hit the 2,048 "
  "token output limit (the longest answer was 260 tokens).")

H2("4.4 Judge and RAGAS Faithfulness")
P("Faithfulness is computed by a separate judge, Qwen/Qwen2.5-72B-Instruct in 4-bit AWQ quantisation [13], [14], served "
  "with vLLM on one 48 GB GPU, greedy decoding, pinned revision. The judge comes from a different model family from the "
  "generator, so the generator is never graded by itself. The judge follows the two RAGAS steps [3]: it splits the "
  "explanation into self-contained statements (using the system description as the question), then decides for each "
  "statement whether it can be inferred from the five passages the generator actually saw. Faithfulness is the share of "
  "supported statements. The quantised 72B model was used because the full-precision model does not fit on the "
  "available hardware; this is a small deviation from the plan and is recorded in the run manifest.")

H2("4.5 Correctness")
P("Correctness is a pure lookup of the predicted label against the Commission label, with no model involved. Accuracy, "
  "precision, recall and F1 for the high-risk class, and macro F1 are reported, broken down by Annex III area and by "
  "filter versus normal cases. A reply that cannot be parsed would count as wrong; none occurred.")

H2("4.6 The Four-Bucket Analysis")
P("Crossing the two axes gives four buckets: right and faithful, right and unfaithful, wrong and faithful, wrong and "
  "unfaithful. The two off-diagonal buckets are the findings. This is a cross-tabulation of two different measures, not "
  "a confusion matrix. An answer counts as faithful at a score of 0.5 or above; because this cut is a reporting choice, "
  "a threshold sweep is also reported. Threshold-free statistics are reported as well: the mean faithfulness of right "
  "and wrong answers, the point-biserial correlation between correctness and faithfulness, and the AUC of faithfulness "
  "as a detector of correct answers (0.5 means no better than chance). 95% confidence intervals come from 2,000 "
  "bootstrap resamples of the 215 inputs. Paired differences in accuracy between systems are tested with McNemar's "
  "exact test.")

H2("4.7 Reproducibility")
table([["Component", "Setting"],
       ["Ground truth", "Commission Annex III draft guidelines, May 2026 (frozen)"],
       ["Test set", "215 inputs, redacted descriptions, labels unchanged"],
       ["Corpus", "400 structural chunks, worked examples held out"],
       ["Embedding", "BAAI/bge-base-en-v1.5, 768 dimensions"],
       ["Retrieval", "Hybrid BM25 + dense, alpha 0.5, top 5 (top 10 logged)"],
       ["Generator", "microsoft/phi-4, vLLM, temperature 0, max 2,048 tokens, pinned revision"],
       ["Judge", "Qwen2.5-72B-Instruct-AWQ, vLLM, temperature 0, pinned revision"],
       ["Faithful cut", "0.5 (sweep 0.3 to 1.0 reported)"],
       ["Hardware", "University JupyterLab server, NVIDIA RTX A6000 48 GB"]],
      [1.5, 4.3], "Table 3. Frozen settings. Model revisions are recorded in the run manifest.")
P("Every run writes, per input, the prompt configuration, the retrieved passages and their scores, the raw model "
  "output, the parsed label and explanation, and the token usage. The whole pipeline runs from one notebook, and "
  "generation resumes from the last completed input if a session is interrupted.")

# ------------------------------------------------------------------ 5
H1("5. Experiments and Results")
P("The results are presented in the order of the pipeline: first whether retrieval finds the right part of the guidelines, then how often each system gets the label right, then how much the label leak had distorted the first run, and finally the central question of whether faithfulness tells right answers from wrong ones. The section ends with a close reading of the answers that are faithful but wrong.")
H2("5.1 Retrieval Quality")
P("Retrieval is scored against the guideline section named as the source of each test input. On the redacted inputs "
  "the correct section is among the top 5 passages for 79.5% of inputs (Table 4). Redaction lowered this only slightly, "
  "even though the inputs became about 30% shorter. Retrieval is weakest for law enforcement (0.565) and migration "
  "(0.652) and better for filter cases than before (0.732).")
ro, rr = L["retrieval"], RD["retrieval"]
table([["Metric", "Original inputs", "Redacted inputs"],
       ["Hit@5 (primary)", f3(ro["overall"]["hit@5"]), f3(rr["overall"]["hit@5"])],
       ["Hit@10", f3(ro["overall"]["hit@10"]), f3(rr["overall"]["hit@10"])],
       ["MRR", f3(ro["overall"]["mrr"]), f3(rr["overall"]["mrr"])],
       ["Hit@5, filter cases (41)", f3(ro["by_edge"]["article-6-3-filter"]["hit@5"]), f3(rr["by_edge"]["article-6-3-filter"]["hit@5"])],
       ["Hit@5, normal cases (174)", f3(ro["by_edge"]["none"]["hit@5"]), f3(rr["by_edge"]["none"]["hit@5"])]],
      [2.4, 1.4, 1.4], "Table 4. Retrieval quality against the gold guideline section.")
figure(FIG + "fig_retrieval_area.png", 5.4, "Figure 4. Retrieval Hit@5 by Annex III area, original against redacted inputs.")

H2("5.2 Correctness of the Three Systems")
rows = [["System", "Accuracy", "HR precision", "HR recall", "HR F1", "Macro F1", "False HR", "Missed HR"]]
mf1 = {B1: 0.682, B2: 0.862, AG: 0.839}
conf = {B1: (60, 7), B2: (17, 9), AG: (22, 9)}
for c, name in [(B1, "Baseline 1: plain LLM"), (B2, "Baseline 2: standard RAG"), (AG, "Agent: structured RAG")]:
    d = RD[c]
    rows.append([name, f3(d["acc"]), f3(d["hr"]["precision"]), f3(d["hr"]["recall"]), f3(d["hr"]["f1"]), f3(mf1[c]),
                 conf[c][0], conf[c][1]])
table(rows, [1.7, 0.7, 0.75, 0.65, 0.6, 0.65, 0.55, 0.6],
      "Table 5. Correctness on the 215 redacted inputs (HR = high-risk, the positive class).")
figure(FIG + "fig_confusion.png", 6.0, "Figure 5. Confusion matrices on the redacted inputs (HR = high-risk).")
P("Retrieval makes a large and significant difference. Both retrieval systems beat the plain model (standard RAG 0.879 "
  "and agent 0.856 against 0.688; McNemar p < 0.001 in both cases). Without retrieval the model labels many systems "
  "high-risk that are not: 60 false high-risk calls against 17 and 22. The structured agent, however, is not better "
  "than standard RAG. It is slightly worse (5 inputs only the agent got right, 10 only standard RAG got right; "
  "McNemar p = 0.30), and it makes more false high-risk calls. Recall is the same for both (0.864).")
P("Figure 6 breaks accuracy down by area. Critical infrastructure is the hardest area for both retrieval systems (0.61 and 0.56), because the test there hinges on whether a system is a safety component rather than on the sector it works in. Without retrieval, law enforcement collapses to 0.26, because the plain model labels most police uses high-risk although none of the law-enforcement examples is; both retrieval systems reach 0.87 or more there.")
figure(FIG + "fig_area_heatmap.png", 6.0, "Figure 6. Accuracy by Annex III area and system on the redacted inputs.")

H2("5.3 Effect of the Label Leak")
figure(FIG + "fig2_accuracy.png", 5.2, "Figure 7. Accuracy before and after removing the verdict text from the inputs.")
P("Figure 7 shows how much the leak inflated the first run. Accuracy fell for every system, most for the plain model "
  "(0.833 to 0.688) and least for standard RAG (0.935 to 0.879). The ranking of the two retrieval systems also flipped: "
  "on the leaky inputs the agent looked best (0.958), on the clean inputs standard RAG is ahead. Without the fix, the "
  "study would have reported both an inflated baseline and the wrong winner.")

H2("5.4 Faithfulness Versus Correctness")
P("Table 6 and Figures 8 and 9 give the main result. For the agent, faithfulness carries no information about whether the "
  "label is right. Wrong answers are on average slightly more faithful than right ones (0.487 against 0.450), the "
  "correlation is close to zero (-0.043, 95% CI -0.17 to 0.08), and as a detector of correct answers faithfulness does "
  "no better than chance (AUC 0.465, 95% CI 0.36 to 0.57). For standard RAG there is a weak signal (AUC 0.615, 95% CI "
  "0.52 to 0.71; correlation 0.142), which is far too weak to be used as a check on legal correctness.")
ag, b2 = RD[AG], RD[B2]
table([["Measure", "Baseline 2: standard RAG", "Agent: structured RAG"],
       ["Mean faithfulness", f3(b2["faith_mean"]), f3(ag["faith_mean"])],
       ["Mean faithfulness, label right", f3(b2["faith_right"]), f3(ag["faith_right"])],
       ["Mean faithfulness, label wrong", f3(b2["faith_wrong"]), f3(ag["faith_wrong"])],
       ["AUC (95% CI)", f"{f3(b2['auc'])} ({f2(b2['auc_ci'][0])} to {f2(b2['auc_ci'][1])})", f"{f3(ag['auc'])} ({f2(ag['auc_ci'][0])} to {f2(ag['auc_ci'][1])})"],
       ["Point-biserial r (95% CI)", f"{f3(b2['r'])} ({f2(b2['r_ci'][0])} to {f2(b2['r_ci'][1])})", f"{f3(ag['r'])} ({f2(ag['r_ci'][0])} to {f2(ag['r_ci'][1])})"],
       ["Right and faithful", 110, 91], ["Right and unfaithful", 79, 93],
       ["Wrong and faithful (faithful but wrong)", 14, 17], ["Wrong and unfaithful", 12, 14]],
      [2.6, 1.7, 1.7], "Table 6. Faithfulness against correctness on the redacted inputs; bucket counts at a cut of 0.5.")
figure(FIG + "fig_buckets.png", 5.6, "Figure 8. The four outcome buckets for the two retrieval systems at a faithful cut of 0.5. The shaded off-diagonal cells are the findings.")
figure(FIG + "fig3_faithfulness.png", 6.0, "Figure 9. Faithfulness of every answer, split by whether the label was right. "
       "Bars mark the group means; the dashed line is the 0.5 cut.")
P("Both off-diagonal buckets are well populated. **Faithful but wrong**: 17 of the agent's 31 wrong answers (55%) and 14 "
  "of standard RAG's 26 score as faithful. **Right but unfaithful**: about half of all correct answers score below 0.5 "
  "(93 of 184 for the agent, 79 of 189 for standard RAG). Part of the second effect is mechanical: the judge marks "
  "sentences that restate the system description as unsupported, because the passages do not contain them. This is how "
  "RAGAS faithfulness is defined and was not changed during the study.")
table([["Faithful cut", "0.3", "0.5", "0.7", "0.9"],
       ["Agent: faithful but wrong", 22, 17, 7, 1], ["Standard RAG: faithful but wrong", 21, 14, 1, 0]],
      [2.4, 0.6, 0.6, 0.6, 0.6], "Table 7. Faithful-but-wrong answers at different cuts. Seven agent errors score 0.7 or higher.")
figure(FIG + "fig_sweep.png", 5.0, "Figure 10. Share of wrong answers that still pass as faithful, as the faithful cut is raised.")
P("The count depends on the cut, as expected (Figure 10), but the agent keeps 7 wrong answers at 0.7 and one wrong answer with a "
  "perfect score of 1.0, so the effect does not rest on a lenient threshold. On the leaky inputs the same analysis gave "
  "an AUC of 0.507 for the agent, but with only 9 wrong answers, which is why the redacted run is the one reported.")

H2("5.5 Article 6(3) Filter Cases")
P("All 41 filter cases have the gold label not-high-risk: the system is in an Annex III area but one of the Article 6(3) "
  "conditions takes it out. Standard RAG handles these best (0.878), the agent and the plain model both reach 0.780. On "
  "normal cases the two retrieval systems are close (0.879 and 0.874) and the plain model is far behind (0.667). The "
  "step-by-step prompt therefore did not help with the one part of the test it was written for; Section 5.6 shows the "
  "typical failure.")
table([["System", "Filter cases (41)", "Normal cases (174)"],
       ["Baseline 1: plain LLM", f3(RD[B1]["filter_acc"]), f3(RD[B1]["none_acc"])],
       ["Baseline 2: standard RAG", f3(RD[B2]["filter_acc"]), f3(RD[B2]["none_acc"])],
       ["Agent: structured RAG", f3(RD[AG]["filter_acc"]), f3(RD[AG]["none_acc"])]],
      [2.2, 1.4, 1.4], "Table 8. Accuracy on Article 6(3) filter cases and on normal cases (redacted inputs).")

H2("5.6 Error Analysis of Faithful-but-Wrong Answers")
P("All 17 faithful-but-wrong answers of the agent were read in full, together with their inputs, retrieved passages and "
  "judge verdicts, and sorted by the main cause of the error (Table 9). Fourteen are false high-risk calls and three are "
  "missed high-risk systems.")
table([["Error type", "Count", "Inputs"],
       ["A. Sector taken as use case: in a critical-infrastructure sector, so labelled high-risk without the safety-component test", 6, "031, 034, 041, 044, 049, 050"],
       ["B. Article 6(3) exception not applied although the task is narrow, preparatory or reviewing completed work", 4, "066, 126, 207, 210"],
       ["C. Profiling or sensitive-data rule applied where it does not fit", 3, "022, 023, 177"],
       ["D. Legal condition misread or a fact invented", 2, "008, 051"],
       ["E. Too cautious on a short description (missed high-risk)", 2, "035, 042"]],
      [3.6, 0.6, 1.6], "Table 9. Causes of the agent's 17 faithful-but-wrong answers (input IDs without the anx3- prefix).")
figure(FIG + "fig_errors.png", 5.2, "Figure 11. Faithful-but-wrong answers of the agent by error type.")
P("Two observations explain why these answers still look faithful. First, retrieval was not the problem: for 14 of the "
  "17 inputs the correct guideline section was among the five passages. The model had the right rule and applied it "
  "wrongly. Second, the judge accepted most of the legal conclusions: 21 of the 31 statements that state a verdict or "
  "reject an exception were marked as supported, because the passages state the general rule correctly. Faithfulness "
  "checks whether a statement can be traced to the passages; it never checks whether the rule was applied correctly to "
  "this system.")
P("The clearest case is input 207, a system that pre-classifies incoming court cases by type. The agent correctly "
  "places it near point 8(a) and then states that no Article 6(3) condition applies, while the guidelines treat it as a "
  "narrow procedural or preparatory task. All six statements in the explanation were judged supported, so the answer "
  "scores a perfect 1.00 and is wrong. Error type A is similar: grid forecasting or network ticketing tools are "
  "labelled high-risk because the passages say that critical infrastructure is an Annex III area, while the passages "
  "also say that only safety components with a direct link to physical harm qualify. Error type C appears to be partly "
  "prompted: the agent's instructions stress that profiling always makes a system high-risk, and the model then finds "
  "profiling where there is none, for example in an anti-money-laundering transaction monitor or in keystroke-based "
  "gender inference, where it also wrongly treats gender as a sensitive attribute under Article 9(1) GDPR.")
P("Two gold labels deserve a caveat. Input 022 comes from an example whose source text is written as a conditional, and "
  "input 035 (a real-time translation tool for public transport staff) is labelled high-risk by the Commission in a way "
  "that is open to debate. Both are kept as labelled; excluding them still leaves 15 faithful-but-wrong answers.")

# ------------------------------------------------------------------ 6
H1("6. Discussion")
P("The results answer the study question clearly. On a real legal classification task with an authoritative answer "
  "key, a high faithfulness score says almost nothing about whether the answer is right. For the structured agent the "
  "two axes are effectively independent, and more than half of its legal errors would pass a faithfulness check. For "
  "standard RAG the link is weak. This is not a failure of the judge: it is what the metric is designed to measure. "
  "Faithfulness asks whether the explanation stays inside the retrieved text, and a wrong legal conclusion can be built "
  "entirely from correctly quoted rules. The error analysis shows that this is exactly what happens: most errors come "
  "from applying a correctly retrieved rule to the wrong facts or skipping one of its conditions.")
P("The comparison of systems adds two practical points. Retrieval matters a lot: the plain model over-predicts "
  "high-risk heavily, and grounding in the guidelines removes most of those false alarms. A more elaborate prompt, "
  "however, did not help. The step-by-step agent was not more accurate than a generic RAG prompt and was worse on the "
  "filter cases it was designed for, and its instruction about profiling seems to have encouraged over-calls. A plausible "
  "explanation is that a procedure written into the prompt invites the model to produce the right-sounding steps without "
  "doing the underlying judgement. This is a hypothesis for further work, not a tested result.")
P("For practitioners the consequence is simple. Faithfulness is necessary but not sufficient for legal compliance "
  "tools. Where a ground truth exists, it should be reported next to a correctness measure, and a high faithfulness "
  "score should not be read as evidence that a classification is right. The label leak found in this study also shows "
  "how easily a benchmark built from worked examples can overstate performance: the leak inflated accuracy by up to 15 "
  "points and reversed the ranking of the two best systems.")

# ------------------------------------------------------------------ 7
H1("7. Conclusion and Future Work")
P("This case study built an evaluation framework for RAG-based EU AI Act high-risk classification with an "
  "authoritative answer key, compared three systems on one shared pipeline, and measured correctness and RAGAS "
  "faithfulness as two independent axes. The main finding is that the two axes come apart: for the structured agent, "
  "faithfulness does not distinguish right from wrong answers (AUC 0.465), and 17 of its 31 errors are faithful but "
  "wrong. The errors are mostly failures to apply a correctly retrieved rule, which faithfulness cannot see. Retrieval "
  "itself improves accuracy substantially, while the structured prompt does not.")
P("**Limitations.** The study uses one generator and one judge, so the result may depend on these models; phi-4 is "
  "smaller than the strongest commercial models, and a stronger generator would make fewer errors but would not change "
  "what faithfulness measures. There is a single annotator and no inter-annotator agreement; the author both reviewed "
  "the redaction and categorised the errors. About 20 redacted inputs still contain wording that describes the legal "
  "test without stating the label. The number of wrong answers is modest (26 and 31), so the confidence intervals are "
  "wide. The test inputs are the Commission's own examples, which are cleaner than real product descriptions, and the "
  "class balance (31% high-risk) and the absence of high-risk examples in three areas are properties of the source. "
  "The faithful cut of 0.5 is a reporting choice, which is why the sweep and threshold-free measures are given.")
P("**Future work.** The most direct extension is a second, stronger generator, including a closed commercial model, "
  "run on the same redacted inputs, to test whether the gap holds across model families and sizes. A second annotator "
  "for the error categories and a small set of realistic product descriptions labelled against the guidelines would "
  "make the result more robust. Finally, a correctness-aware check that tests whether the conclusion follows from the "
  "retrieved rule and the stated facts could be evaluated against the same answer key; fixing the metric was out of "
  "scope here, but the framework and the labelled failure cases make that a natural next step, and a possible basis for "
  "a thesis or paper.")

# ------------------------------------------------------------------ 8
H1("8. Code and Data Availability")
P("The source code (corpus builder, retriever, generation harness, judge, scoring, redaction scripts and the run "
  "notebooks) is available at:")
P("https://github.com/rahulrawat20022002/case_study2", align=WD_ALIGN_PARAGRAPH.LEFT)
P("The repository contains the original and redacted test sets, the redaction log with the author's review, the "
  "retrieval corpus, and the run manifests with pinned model revisions. Model weights are not included and are "
  "downloaded from Hugging Face at the recorded revisions.")

# ------------------------------------------------------------------ 9
H1("9. References")
refs = [
    "Regulation (EU) 2024/1689 of the European Parliament and of the Council of 13 June 2024 laying down harmonised rules on artificial intelligence (Artificial Intelligence Act), Official Journal of the European Union, L series, 12 July 2024.",
    "European Commission, \"Draft Guidelines on the classification of high-risk AI systems: Annex III,\" draft, May 2026.",
    "S. Es, J. James, L. Espinosa-Anke, S. Schockaert, \"RAGAS: Automated Evaluation of Retrieval Augmented Generation,\" in Proc. EACL 2024: System Demonstrations, pp. 150-158.",
    "P. Lewis et al., \"Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks,\" in Advances in Neural Information Processing Systems (NeurIPS) 33, 2020.",
    "J. Saad-Falcon, O. Khattab, C. Potts, M. Zaharia, \"ARES: An Automated Evaluation Framework for Retrieval-Augmented Generation Systems,\" in Proc. NAACL 2024, pp. 338-354.",
    "I. Chalkidis, A. Jana, D. Hartung, M. Bommarito, I. Androutsopoulos, D. M. Katz, N. Aletras, \"LexGLUE: A Benchmark Dataset for Legal Language Understanding in English,\" in Proc. ACL 2022, pp. 4310-4330.",
    "OECD, \"OECD Framework for the Classification of AI Systems,\" OECD Digital Economy Papers, No. 323, OECD Publishing, Paris, 2022.",
    "National Institute of Standards and Technology, \"Artificial Intelligence Risk Management Framework (AI RMF 1.0),\" NIST AI 100-1, 2023.",
    "S. Xiao, Z. Liu, P. Zhang, N. Muennighoff, D. Lian, J.-Y. Nie, \"C-Pack: Packed Resources For General Chinese Embeddings,\" in Proc. SIGIR 2024. arXiv:2309.07597.",
    "S. Robertson, H. Zaragoza, \"The Probabilistic Relevance Framework: BM25 and Beyond,\" Foundations and Trends in Information Retrieval, vol. 3, no. 4, pp. 333-389, 2009.",
    "M. Abdin et al., \"Phi-4 Technical Report,\" arXiv:2412.08905, 2024.",
    "W. Kwon et al., \"Efficient Memory Management for Large Language Model Serving with PagedAttention,\" in Proc. SOSP 2023.",
    "Qwen Team, \"Qwen2.5 Technical Report,\" arXiv:2412.15115, 2024.",
    "J. Lin et al., \"AWQ: Activation-aware Weight Quantization for On-Device LLM Compression and Acceleration,\" in Proc. MLSys 2024.",
]
for i, r in enumerate(refs, 1):
    P(f"[{i}] {r}", align=WD_ALIGN_PARAGRAPH.LEFT, size=10)

import sys, os
pages = json.load(open(sys.argv[1])) if len(sys.argv) > 1 and os.path.exists(sys.argv[1]) else {}
anchor = TOC_ANCHOR[0]
for lvl, t in [(1, "Abstract"), (1, "Contents")] + [h for h in HEADINGS if h[1] not in ("Abstract", "Contents")]:
    p = doc.add_paragraph(style="toc 1" if lvl == 1 else "toc 2")
    pf = p.paragraph_format; pf.space_after = Pt(2); pf.space_before = Pt(0)
    pf.tab_stops.add_tab_stop(Emu(8640 * 635), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.add_run(f"{t}\t{pages.get(t, 0)}")
    anchor._p.addprevious(p._p)
anchor._p.getparent().remove(anchor._p)
json.dump(HEADINGS, open("/tmp/claude-0/-home-claude/3f0a219a-3c0a-53e0-8360-748fb7d57569/scratchpad/headings.json", "w"))
doc.save(OUT)
print("saved", OUT)
